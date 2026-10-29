"""Build the editable LAMS report from its Markdown source on Windows.

Requires python-docx and Pillow. Export the resulting DOCX to PDF in Word,
then inspect every PDF page before committing either artifact.
"""

from __future__ import annotations

import re
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'ho-so-nop-bai/bao-cao/Bao-cao-tieu-luan-LAMS.md'
OUT = ROOT / 'ho-so-nop-bai/bao-cao/Bao-cao-tieu-luan-LAMS.docx'
NAVY = RGBColor(27, 54, 94)
GRAY = RGBColor(86, 94, 106)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2)
section.bottom_margin = Cm(2)
section.left_margin = Cm(2.6)
section.right_margin = Cm(2)
section.header.paragraphs[0].text = ''

for name, size, bold in [('Normal', 13, False), ('Title', 22, True), ('Heading 1', 16, True), ('Heading 2', 14, True), ('Heading 3', 13, True)]:
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(30, 30, 30)
    style.paragraph_format.space_after = Pt(3 if name == 'Normal' else 8)
    if name == 'Normal': style.paragraph_format.line_spacing = 1.3
    else: style.paragraph_format.keep_with_next = True

doc.styles['Caption'].font.name = 'Times New Roman'
doc.styles['Caption'].font.size = Pt(10.5)
doc.styles['Caption'].font.color.rgb = RGBColor(45, 45, 45)
doc.styles['Normal'].paragraph_format.widow_control = True
for name in ('TOC 1', 'TOC 2'):
    style = doc.styles.add_style(name, 1) if name not in doc.styles else doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11.5)
    style.paragraph_format.line_spacing = 1.05
    style.paragraph_format.space_after = Pt(2)

def page_orientation(landscape=False):
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Cm(29.7 if landscape else 21)
    section.page_height = Cm(21 if landscape else 29.7)
    section.top_margin = Cm(1.4 if landscape else 2)
    section.bottom_margin = Cm(1.4 if landscape else 2)
    section.left_margin = Cm(1.6 if landscape else 2.6)
    section.right_margin = Cm(1.6 if landscape else 2)
    # The section-break paragraph must not spill alone after a full table.
    boundary = doc.paragraphs[-1]
    boundary.paragraph_format.space_before = Pt(0)
    boundary.paragraph_format.space_after = Pt(0)
    boundary.paragraph_format.line_spacing = Pt(1)
    boundary.paragraph_format.keep_with_next = False
    boundary.add_run().font.size = Pt(1)
    # Section changes must continue page numbering after the cover.
    for element in section._sectPr.findall(qn('w:pgNumType')):
        section._sectPr.remove(element)

def add_rich(paragraph, text: str):
    chunks = re.split(r'(\[[^]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`)', text)
    for chunk in chunks:
        link = re.fullmatch(r'\[([^]]+)\]\(([^)]+)\)', chunk)
        if link:
            target = link.group(2)
            h = OxmlElement('w:hyperlink')
            if target.startswith('#'):
                h.set(qn('w:anchor'), target[1:])
            else:
                rid = paragraph.part.relate_to(target,
                    'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
                    is_external=True)
                h.set(qn('r:id'), rid)
            run = OxmlElement('w:r')
            t = OxmlElement('w:t'); t.text = link.group(1); run.append(t); h.append(run)
            paragraph._p.append(h)
        elif chunk.startswith('**') and chunk.endswith('**'):
            paragraph.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith('`') and chunk.endswith('`'):
            run = paragraph.add_run(chunk[1:-1]); run.font.name = 'Consolas'; run.font.size = Pt(10)
        else:
            paragraph.add_run(chunk)

def table_from_lines(rows: list[str]):
    data = [[c.strip() for c in row.strip().strip('|').split('|')] for row in rows]
    data = [row for row in data if not all(re.fullmatch(r':?-{3,}:?', c) for c in row)]
    table = doc.add_table(rows=0, cols=len(data[0]))
    table.autofit = False
    if data[0] == ['Nhóm', 'Bảng đang có trong migration', 'Vai trò']: widths = [3.1, 8.7, 4.4]
    elif data[0][0] == 'Mã / UC / ưu tiên': widths = [2.7, 6.5, 7.0]
    elif len(data[0]) == 3 and data[0][0] == 'Mã': widths = [2.0, 8.0, 6.2]
    elif len(data[0]) == 3: widths = [4.5, 6.5, 5.2]
    elif len(data[0]) >= 6:
        widths = {
            6: [3.0, 3.8, 4.2, 4.5, 3.9, 5.7],
            7: [1.2, 2.5, 4.0, 4.3, 3.6, 4.9, 4.6],
            8: [1.2, 3.7, 1.8, 1.8, 4.2, 3.7, 4.2, 4.5],
        }[len(data[0])]
        widths = [width * 26.5 / sum(widths) for width in widths]
    else: widths = [16.2 / len(data[0])] * len(data[0])
    for j, width in enumerate(widths): table.columns[j].width = Cm(width)
    for index, values in enumerate(data):
        cells = table.add_row().cells
        for j, value in enumerate(values):
            cells[j].width = Cm(widths[j])
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[j].text = ''
            add_rich(cells[j].paragraphs[0], value)
            for p in cells[j].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 0 and data[0][0] == 'Mã' else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.left_indent = Pt(0)
                p.paragraph_format.space_after = Pt(1 if len(data[0]) >= 6 else 3)
                p.paragraph_format.line_spacing = 1.05
                p.paragraph_format.keep_with_next = index == 0
                p.paragraph_format.keep_together = True
                for run in p.runs: run.font.size = Pt(10.5)
            tcPr = cells[j]._tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), '1B365E' if index == 0 else ('F1F5FA' if index % 2 == 0 else 'FFFFFF'))
            tcPr.append(shd)
            mar = OxmlElement('w:tcMar')
            for side in ('top', 'left', 'bottom', 'right'):
                padding = '60' if len(data[0]) >= 6 and side in ('top', 'bottom') else '100'
                el = OxmlElement('w:' + side); el.set(qn('w:w'), padding); el.set(qn('w:type'), 'dxa'); mar.append(el)
            tcPr.append(mar)
            for run in cells[j].paragraphs[0].runs:
                if index == 0: run.font.bold = True; run.font.color.rgb = RGBColor(255, 255, 255)
        if index == 0:
            trPr = table.rows[-1]._tr.get_or_add_trPr()
            repeat = OxmlElement('w:tblHeader'); repeat.set(qn('w:val'), 'true'); trPr.append(repeat)
        trPr = table.rows[-1]._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit'); trPr.append(no_split)
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + side); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D9D9D9'); borders.append(e)
    tblPr.append(borders)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = Pt(1)
    p.add_run().font.size = Pt(1)

def cover_line(content, size, bold=False, before=0, after=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(content)
    run.font.name = 'Times New Roman'; run.font.size = Pt(size); run.bold = bold
    run.font.color.rgb = RGBColor(27, 54, 94) if bold else RGBColor(40, 45, 52)
    return p

cover_line('ĐẠI HỌC PHENIKAA', 18, True, before=8, after=3)
cover_line('TRƯỜNG CÔNG NGHỆ THÔNG TIN', 18, True, after=4)
cover_line('HỌC PHẦN: THIẾT KẾ WEB NÂNG CAO', 14, after=3)
cover_line('-----***-----', 14, after=10)
logo = doc.add_paragraph()
logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
logo.paragraph_format.space_before = Pt(10)
logo.paragraph_format.space_after = Pt(16)
logo.add_run().add_picture(str(SOURCE.parent / 'hinh/logo-phenikaa.png'), width=Cm(8.2))
band = cover_line('BÁO CÁO BÀI TẬP LỚN', 22, True, after=8)
shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), '1B365E')
band._p.get_or_add_pPr().append(shade)
band.runs[0].font.color.rgb = RGBColor(255,255,255)
title = cover_line('HỆ THỐNG QUẢN LÍ VÀ KHAI THÁC THƯ VIỆN', 17, True, after=4)
borders=OxmlElement('w:pBdr'); bottom=OxmlElement('w:bottom')
bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'12'); bottom.set(qn('w:color'),'F58220')
borders.append(bottom); title._p.get_or_add_pPr().append(borders)
cover_line('LAMS — Báo cáo cá nhân', 12, after=12)
members = doc.add_table(rows=0, cols=2)
members.alignment = 1
members.autofit = False
members.columns[0].width = Cm(4)
members.columns[1].width = Cm(9.8)
for index, (label, value) in enumerate([
    ('Thông tin', 'Nội dung'),
    ('Tác giả báo cáo', 'Nguyễn Trọng Tuấn — 23010690'),
    ('Lớp học phần', 'Thiết kế Web nâng cao-1-1-26(N01)'),
    ('Giảng viên', ''),
    ('Nhóm thực hiện', 'Nguyễn Trọng Tuấn — 23010690\nVũ Đức Toàn — 23010062'),
    ('Phạm vi báo cáo', 'Dự án chung của nhóm và đóng góp cá nhân của Nguyễn Trọng Tuấn'),
]):
    row = members.add_row()
    for j, value in enumerate((label,value)):
        cell=row.cells[j]; cell.text=value
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'214F77' if index==0 else ('DCE6F1' if j==0 else 'F4F7FB'))
        cell._tc.get_or_add_tcPr().append(sh)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after=Pt(5)
            paragraph.paragraph_format.space_before=Pt(3)
            paragraph.paragraph_format.line_spacing=1.05
            for run in paragraph.runs:
                run.font.size=Pt(11); run.bold=(j==0 or index==0)
                run.font.color.rgb=RGBColor(255,255,255) if index==0 else RGBColor(30,30,30)
cover_line('Hà Nội, năm học 2026 - 2027', 12, before=24)

section = doc.add_section(WD_SECTION_START.NEW_PAGE)
section.footer.is_linked_to_previous = False
section.header.is_linked_to_previous = False
section.header.paragraphs[0].text = ''
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('Trang ')
begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
code = OxmlElement('w:instrText'); code.set(qn('xml:space'), 'preserve'); code.text = ' PAGE '
sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate')
end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
footer.add_run()._r.extend([begin, code, sep, end])
page_numbering = OxmlElement('w:pgNumType'); page_numbering.set(qn('w:start'), '1')
section._sectPr.append(page_numbering)

lines = SOURCE.read_text(encoding='utf-8').splitlines()
figure_titles = [re.match(r'!\[(Hình \d+\. [^]]+)\]', s).group(1)
                 for s in lines if re.match(r'!\[(Hình \d+\. [^]]+)\]', s)]
table_titles = [s for s in lines if re.match(r'Bảng \d+\. ', s)]

def add_field(paragraph, instruction, result='…'):
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    code = OxmlElement('w:instrText'); code.set(qn('xml:space'), 'preserve'); code.text = ' ' + instruction + ' '
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate')
    shown = OxmlElement('w:t'); shown.text = result
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    paragraph.add_run()._r.extend([begin, code, sep, shown, end])

def add_bookmark(paragraph, name):
    number = len(doc.element.body.xpath('.//w:bookmarkStart')) + 1
    start = OxmlElement('w:bookmarkStart'); start.set(qn('w:id'), str(number)); start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(number))
    paragraph._p.insert(0, start); paragraph._p.append(end)

def caption(value, kind):
    p = doc.add_paragraph(style='Caption')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if kind == 'fig' else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.keep_with_next = kind == 'tab'
    match = re.match(r'(Hình|Bảng) (\d+)(.*)', value)
    p.add_run(match.group(1) + ' ')
    add_field(p, f'SEQ {"Figure" if kind == "fig" else "Table"} \\* ARABIC', match.group(2))
    add_rich(p, match.group(3))
    add_bookmark(p, f'{"fig" if kind == "fig" else "tbl"}-{match.group(2)}')

def list_items(items, kind):
    for title in items:
        number = re.search(r'\d+', title).group()
        p = doc.add_paragraph(style='Normal')
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.space_after = Pt(5)
        add_rich(p, title + '  ')
        for run in p.runs: run.font.size = Pt(11)
        add_field(p, f'PAGEREF {"fig" if kind == "fig" else "tbl"}-{number} \\h')

table_rows = []
in_use_case_appendix = False
wide_table = False
def flush_table():
    global table_rows, wide_table
    if table_rows:
        table_from_lines(table_rows)
        table_rows = []
        if wide_table:
            page_orientation(False)
            wide_table = False
for line_index, line in enumerate(lines):
    stripped = line.strip()
    if stripped.startswith('|'): table_rows.append(stripped); continue
    flush_table()
    if not stripped: continue
    if stripped == '<!--TOC-->':
        add_field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u', 'Cập nhật mục lục trong Word')
        continue
    if stripped == '<!--LISTFIG-->':
        list_items(figure_titles, 'fig'); continue
    if stripped == '<!--LISTTAB-->':
        list_items(table_titles, 'tab'); continue
    if stripped == '<!--FIGUREPAGE-->':
        doc.add_page_break(); continue
    if stripped.startswith('!['):
        match = re.match(r'!\[([^]]+)\]\(([^)]+)\)', stripped)
        if match:
            path = (SOURCE.parent / match.group(2)).resolve()
            if not path.is_file(): raise FileNotFoundError(path)
            with Image.open(path) as img: width, height = img.size
            diagram = path.name in {'use-cases.png', 'use-cases-staff.png', 'class-diagram.png', 'class-catalog.png', 'class-inventory.png',
                                   'erd-core.png', 'erd-insight.png', 'erd-identity.png', 'erd-inventory.png',
                                   'hold-sequence.png', 'inventory-flow.png'}
            max_height = 10.5 if path.name == 'use-cases.png' else (19 if diagram else 14)
            picture_width = min(16.2, max_height * width / height)
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(path), width=Cm(picture_width))
            caption(match.group(1), 'fig')
        continue
    if re.match(r'Bảng \d+\. ', stripped):
        upcoming = next((s.strip() for s in lines[line_index+1:] if s.strip()), '')
        if upcoming.startswith('|') and len(upcoming.strip('|').split('|')) >= 6:
            page_orientation(True)
            wide_table = True
        caption(stripped, 'tab'); continue
    if stripped.startswith('<!--'):
        continue
    if stripped.startswith('# '):
        p = doc.add_paragraph(stripped[2:], 'Heading 1')
        if stripped[2:] in {'Mục lục', 'Danh mục hình', 'Danh mục bảng', 'Từ viết tắt'} or stripped.startswith('# Chương 1.'):
            p.paragraph_format.page_break_before = True
        if not stripped.startswith(('# Chương', '# Tài liệu tham khảo', '# Phụ lục')):
            p.style = 'Title' if stripped.startswith('# Báo cáo') else 'Subtitle'
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(16)
                run.bold = True
    elif stripped.startswith('## '):
        if stripped.startswith('## Phụ lục E.'): in_use_case_appendix = True
        doc.add_paragraph(stripped[3:], 'Heading 2')
    elif stripped.startswith('### '):
        doc.add_paragraph(stripped[4:], 'Heading 3')
    else:
        p = doc.add_paragraph(style='Normal')
        add_rich(p, stripped)
        if in_use_case_appendix:
            p.paragraph_format.line_spacing = 1.1
            p.paragraph_format.space_after = Pt(1)
            for run in p.runs:
                if run.font.name != 'Consolas': run.font.size = Pt(12)
flush_table()
doc.core_properties.title = 'Báo cáo bài tập lớn Hệ thống quản lí và khai thác thư viện LAMS'
doc.core_properties.subject = 'Báo cáo cá nhân của Nguyễn Trọng Tuấn về dự án LAMS'
doc.core_properties.author = 'Nguyễn Trọng Tuấn'
doc.save(OUT)
print(OUT)
